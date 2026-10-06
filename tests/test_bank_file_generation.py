"""Phase 4: bank disbursement batch generation for Approved payroll runs.

Covers: refusal for a non-Approved run, refusal (with the full affected-
employee list, nothing persisted) for missing/invalid bank details or
non-positive net pay, a clean generation's exact column layout/order/
amounts, byte-for-byte re-download via checksum, and version incrementing
across repeated generations for the same run.
"""
import hashlib
from datetime import date
from io import BytesIO

from openpyxl import load_workbook

from app.models.company_bank_config import CompanyBankConfig
from app.models.disbursement import DisbursementBatch, DisbursementBatchItem
from app.models.employee import Employee, EmployeeBankAccount
from app.models.pay_schedule import PaySchedule
from app.models.payroll_run import PayrollRun, PayrollRunLineItem

PERIOD_START = date(2026, 9, 1)
PERIOD_END = date(2026, 9, 30)

EXPECTED_COLUMNS = [
    "Beneficiary Name",
    "Beneficiary Account Number",
    "IFSC Code",
    "Amount",
    "Narration",
    "Debit Account Number",
    "Value Date",
    "Payment Mode",
]


def _make_schedule(db_session) -> PaySchedule:
    schedule = PaySchedule(
        schedule_name="Monthly - All",
        pay_frequency="Monthly",
        payment_day_rule={"type": "LastDayOfMonth"},
        holiday_adjustment="Before",
        is_active=True,
        applies_to=None,
    )
    db_session.add(schedule)
    db_session.commit()
    db_session.refresh(schedule)
    return schedule


def _make_run(db_session, schedule, status="Approved") -> PayrollRun:
    run = PayrollRun(
        pay_schedule_id=schedule.id,
        period_start=PERIOD_START,
        period_end=PERIOD_END,
        pay_date=date(2026, 10, 5),
        run_type="Regular",
        status=status,
        source="InAppCalculation",
        total_gross=0,
        total_deductions=0,
        total_net=0,
        employee_count=0,
    )
    db_session.add(run)
    db_session.commit()
    db_session.refresh(run)
    return run


def _make_employee(db_session, code, bank=None) -> Employee:
    employee = Employee(
        employee_code=code,
        first_name=code,
        status="Active",
        date_of_joining=date(2025, 1, 1),
        basic_salary=50000,
    )
    db_session.add(employee)
    db_session.commit()
    db_session.refresh(employee)

    if bank is not None:
        account = EmployeeBankAccount(
            employee_id=employee.id,
            bank_name=bank.get("bank_name", "HDFC Bank"),
            account_holder_name=bank.get("account_holder_name", f"{code} Holder"),
            account_number=bank["account_number"],
            ifsc_code=bank["ifsc_code"],
            branch_name="Main Branch",
            account_type="Savings",
            is_primary=True,
            is_active=True,
        )
        db_session.add(account)
        db_session.commit()
    return employee


def _make_line_item(db_session, run, employee, net_pay, status="Included") -> PayrollRunLineItem:
    line_item = PayrollRunLineItem(
        payroll_run_id=run.id,
        employee_id=employee.id,
        employee_snapshot={
            "employee_code": employee.employee_code,
            "name": f"{employee.first_name} {employee.last_name or ''}".strip(),
            "status": employee.status,
        },
        gross_earnings=net_pay,
        total_deductions=0,
        total_reimbursements=0,
        net_pay=net_pay,
        status=status,
        computation_breakdown=[],
        error_message=None,
    )
    db_session.add(line_item)
    db_session.commit()
    db_session.refresh(line_item)
    return line_item


def _seed_company_bank_config(db_session) -> CompanyBankConfig:
    config = CompanyBankConfig(
        company_code="STAUNCH",
        debit_account_number="0000000000",
        debit_account_bank_name="ICICI Bank (PLACEHOLDER)",
        is_active=True,
    )
    db_session.add(config)
    db_session.commit()
    db_session.refresh(config)
    return config


def test_generation_refused_if_run_not_approved(client, db_session):
    _seed_company_bank_config(db_session)
    schedule = _make_schedule(db_session)
    run = _make_run(db_session, schedule, status="Calculated")
    emp = _make_employee(db_session, "E-BF1", bank={"account_number": "1234567890", "ifsc_code": "HDFC0001234"})
    _make_line_item(db_session, run, emp, 50000)

    resp = client.post(f"/api/v1/payroll-runs/{run.id}/disbursement-batches", json={})
    assert resp.status_code == 400, resp.text
    assert "Approved" in resp.json()["detail"]
    assert db_session.query(DisbursementBatch).count() == 0


def test_generation_refused_with_affected_employee_list_and_nothing_persisted(client, db_session):
    _seed_company_bank_config(db_session)
    schedule = _make_schedule(db_session)
    run = _make_run(db_session, schedule, status="Approved")

    emp_no_bank = _make_employee(db_session, "E-BF2")  # no bank account at all
    emp_bad_ifsc = _make_employee(
        db_session, "E-BF3", bank={"account_number": "1111111111", "ifsc_code": "BADCODE1"}
    )
    emp_zero_pay = _make_employee(
        db_session, "E-BF4", bank={"account_number": "2222222222", "ifsc_code": "HDFC0001234"}
    )

    _make_line_item(db_session, run, emp_no_bank, 40000)
    _make_line_item(db_session, run, emp_bad_ifsc, 40000)
    _make_line_item(db_session, run, emp_zero_pay, 0)

    resp = client.post(f"/api/v1/payroll-runs/{run.id}/disbursement-batches", json={})
    assert resp.status_code == 422, resp.text
    detail = resp.json()["detail"]
    issues_text = " ".join(detail["issues"])
    assert "E-BF2" in issues_text and "no active bank account" in issues_text
    assert "E-BF3" in issues_text and "invalid IFSC" in issues_text
    assert "E-BF4" in issues_text and "net pay" in issues_text
    assert len(detail["issues"]) == 3

    assert db_session.query(DisbursementBatch).count() == 0
    assert db_session.query(DisbursementBatchItem).count() == 0


def test_successful_generation_has_exact_columns_and_amounts(client, db_session):
    _seed_company_bank_config(db_session)
    schedule = _make_schedule(db_session)
    run = _make_run(db_session, schedule, status="Approved")

    emp1 = _make_employee(
        db_session,
        "E-BF5",
        bank={"account_number": "3333333333", "ifsc_code": "HDFC0001234", "account_holder_name": "E-BF5 Holder"},
    )
    emp2 = _make_employee(
        db_session,
        "E-BF6",
        bank={"account_number": "4444444444", "ifsc_code": "ICIC0005678", "account_holder_name": "E-BF6 Holder"},
    )
    _make_line_item(db_session, run, emp1, 84950.50)
    _make_line_item(db_session, run, emp2, 250000)  # > 2,00,000 -> RTGS

    resp = client.post(f"/api/v1/payroll-runs/{run.id}/disbursement-batches", json={})
    assert resp.status_code == 201, resp.text
    batch_data = resp.json()
    assert batch_data["total_employee_count"] == 2
    assert batch_data["version"] == 1
    assert batch_data["format_key"] == "generic_neft_xlsx_v1"
    download_url = batch_data["download_url"]

    db_batch = db_session.query(DisbursementBatch).filter(DisbursementBatch.id == batch_data["id"]).first()
    assert db_batch is not None
    items = db_session.query(DisbursementBatchItem).filter(DisbursementBatchItem.disbursement_batch_id == db_batch.id).all()
    assert len(items) == 2

    download_resp = client.get(f"/api/v1{download_url}")
    assert download_resp.status_code == 200, download_resp.text
    assert download_resp.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    wb = load_workbook(BytesIO(download_resp.content))
    ws = wb["Salary Disbursement"]
    header_row = [cell.value for cell in ws[1]]
    assert header_row == EXPECTED_COLUMNS
    assert ws[1][0].font.bold is True

    data_rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(data_rows) == 2

    by_account = {row[1]: row for row in data_rows}
    row1 = by_account["3333333333"]
    assert row1[0] == "E-BF5 Holder"
    assert row1[2] == "HDFC0001234"
    assert row1[3] == 84950.50
    assert row1[5] == "0000000000"
    assert row1[7] == "NEFT"

    row2 = by_account["4444444444"]
    assert row2[0] == "E-BF6 Holder"
    assert row2[3] == 250000.0
    assert row2[7] == "RTGS"


def test_redownload_is_byte_for_byte_identical(client, db_session):
    _seed_company_bank_config(db_session)
    schedule = _make_schedule(db_session)
    run = _make_run(db_session, schedule, status="Approved")
    emp = _make_employee(db_session, "E-BF7", bank={"account_number": "5555555555", "ifsc_code": "HDFC0001234"})
    _make_line_item(db_session, run, emp, 60000)

    resp = client.post(f"/api/v1/payroll-runs/{run.id}/disbursement-batches", json={})
    assert resp.status_code == 201, resp.text
    batch_data = resp.json()
    checksum = batch_data["checksum"]

    first_download = client.get(f"/api/v1/disbursement-batches/{batch_data['id']}/download")
    assert first_download.status_code == 200
    first_checksum = hashlib.sha256(first_download.content).hexdigest()
    assert first_checksum == checksum

    second_download = client.get(f"/api/v1/disbursement-batches/{batch_data['id']}/download")
    assert second_download.status_code == 200
    second_checksum = hashlib.sha256(second_download.content).hexdigest()
    assert second_checksum == checksum
    assert second_download.content == first_download.content


def test_generating_twice_increments_version(client, db_session):
    _seed_company_bank_config(db_session)
    schedule = _make_schedule(db_session)
    run = _make_run(db_session, schedule, status="Approved")
    emp = _make_employee(db_session, "E-BF8", bank={"account_number": "6666666666", "ifsc_code": "HDFC0001234"})
    _make_line_item(db_session, run, emp, 70000)

    first_resp = client.post(f"/api/v1/payroll-runs/{run.id}/disbursement-batches", json={})
    assert first_resp.status_code == 201, first_resp.text
    assert first_resp.json()["version"] == 1

    second_resp = client.post(f"/api/v1/payroll-runs/{run.id}/disbursement-batches", json={})
    assert second_resp.status_code == 201, second_resp.text
    assert second_resp.json()["version"] == 2

    list_resp = client.get(f"/api/v1/payroll-runs/{run.id}/disbursement-batches")
    assert list_resp.status_code == 200
    versions = sorted(b["version"] for b in list_resp.json())
    assert versions == [1, 2]

    # Both batches remain independently downloadable.
    for batch in list_resp.json():
        dl = client.get(f"/api/v1{batch['download_url']}")
        assert dl.status_code == 200


def test_disbursement_data_endpoint_requires_approved(client, db_session):
    _seed_company_bank_config(db_session)
    schedule = _make_schedule(db_session)
    run = _make_run(db_session, schedule, status="Draft")
    resp = client.get(f"/api/v1/payroll-runs/{run.id}/disbursement-data")
    assert resp.status_code == 400


def test_disbursement_data_endpoint_returns_rows_for_approved_run(client, db_session):
    _seed_company_bank_config(db_session)
    schedule = _make_schedule(db_session)
    run = _make_run(db_session, schedule, status="Approved")
    emp = _make_employee(db_session, "E-BF9", bank={"account_number": "7777777777", "ifsc_code": "HDFC0001234"})
    _make_line_item(db_session, run, emp, 55000)

    resp = client.get(f"/api/v1/payroll-runs/{run.id}/disbursement-data")
    assert resp.status_code == 200, resp.text
    rows = resp.json()
    assert len(rows) == 1
    assert rows[0]["employee_code"] == "E-BF9"
    assert rows[0]["account_number"] == "7777777777"
    assert rows[0]["net_pay"] == 55000
